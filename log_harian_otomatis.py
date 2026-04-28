#!/usr/bin/env python3
"""
Excel Logbook Editor
Desktop application untuk membaca dan menyimpan file Excel.
Fitur: buka file, import data pasien, proses data otomatis,
       pengisian otomatis no CM, tanggal, dan alat medis EKG.
"""

import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import re
from collections import OrderedDict
from datetime import datetime
import io

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, OneCellAnchor, AnchorMarker
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, OneCellAnchor, AnchorMarker

BULAN_INDO = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April",
    5: "Mei", 6: "Juni", 7: "Juli", 8: "Agustus",
    9: "September", 10: "Oktober", 11: "November", 12: "Desember"
}


class ExcelEditorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel Logbook Editor")
        self.root.geometry("1200x700")
        self.root.minsize(900, 500)

        self.workbook = None
        self.filepath = None
        self.original_filepath = None
        self.current_sheet_name = None
        self.modified = False
        self.patient_data = None

        self._build_ui()

    # ------------------------------------------------------------------ UI
    def _build_ui(self):
        # Menu bar
        menubar = tk.Menu(self.root)
        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="Buka File Excel...", command=self.open_file, accelerator="Ctrl+O")
        file_menu.add_command(label="Simpan", command=self.save_file, accelerator="Ctrl+S")
        file_menu.add_command(label="Simpan Sebagai...", command=self.save_as, accelerator="Ctrl+Shift+S")
        file_menu.add_separator()
        file_menu.add_command(label="Keluar", command=self.on_close)
        menubar.add_cascade(label="File", menu=file_menu)

        tools_menu = tk.Menu(menubar, tearoff=0)
        tools_menu.add_command(label="Load Data Pasien...", command=self.load_data_pasien)
        tools_menu.add_command(label="Proses Data", command=self.proses_data)
        menubar.add_cascade(label="Tools", menu=tools_menu)

        self.root.config(menu=menubar)
        self.root.bind("<Control-o>", lambda e: self.open_file())
        self.root.bind("<Control-s>", lambda e: self.save_file())
        self.root.bind("<Control-S>", lambda e: self.save_as())
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

        # Toolbar
        toolbar = ttk.Frame(self.root)
        toolbar.pack(fill=tk.X, padx=8, pady=(8, 0))

        ttk.Button(toolbar, text="Buka File", command=self.open_file).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Simpan", command=self.save_file).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Simpan Sebagai", command=self.save_as).pack(side=tk.LEFT, padx=2)
        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=8)
        ttk.Button(toolbar, text="Load Data Pasien", command=self.load_data_pasien).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="Proses Data", command=self.proses_data).pack(side=tk.LEFT, padx=2)

        self.file_label = ttk.Label(toolbar, text="Belum ada file dibuka", foreground="gray")
        self.file_label.pack(side=tk.RIGHT, padx=8)

        self.status_label = ttk.Label(toolbar, text="", foreground="green")
        self.status_label.pack(side=tk.RIGHT, padx=8)

        # Sheet tabs
        self.tab_frame = ttk.Frame(self.root)
        self.tab_frame.pack(fill=tk.X, padx=8, pady=(8, 0))

        self.tab_canvas = tk.Canvas(self.tab_frame, height=32, highlightthickness=0)
        self.tab_scrollbar = ttk.Scrollbar(self.tab_frame, orient=tk.HORIZONTAL, command=self.tab_canvas.xview)
        self.tab_inner = ttk.Frame(self.tab_canvas)

        self.tab_canvas.configure(xscrollcommand=self.tab_scrollbar.set)
        self.tab_scrollbar.pack(side=tk.BOTTOM, fill=tk.X)
        self.tab_canvas.pack(side=tk.TOP, fill=tk.X)
        self.tab_canvas.create_window((0, 0), window=self.tab_inner, anchor="nw")
        self.tab_inner.bind("<Configure>", lambda e: self.tab_canvas.configure(scrollregion=self.tab_canvas.bbox("all")))

        # Table area
        table_frame = ttk.Frame(self.root)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        self.tree = ttk.Treeview(table_frame, show="headings", selectmode="browse")

        vsb = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        hsb = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

        # Status bar
        self.statusbar = ttk.Label(self.root, text="Siap", relief=tk.SUNKEN, anchor=tk.W)
        self.statusbar.pack(fill=tk.X, side=tk.BOTTOM)

        # Welcome screen
        self._show_welcome()

    def _show_welcome(self):
        for widget in self.tab_inner.winfo_children():
            widget.destroy()
        self.tree["columns"] = ()
        self.tree.delete(*self.tree.get_children())

    # ------------------------------------------------------------------ FILE OPS
    def open_file(self):
        path = filedialog.askopenfilename(
            title="Pilih File Excel",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if not path:
            return
        try:
            self.workbook = load_workbook(path)
            self.filepath = path
            self.original_filepath = path
            self.modified = False
            self.patient_data = None
            self.file_label.config(text=path.split("/")[-1].split("\\")[-1], foreground="black")
            self._render_tabs()
            self._switch_sheet(self.workbook.sheetnames[0])
            self.statusbar.config(text=f"File dibuka: {path}")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal membuka file:\n{e}")

    def save_file(self):
        if not self.workbook or not self.filepath:
            self.save_as()
            return
        # Never overwrite original file — auto-generate new filename
        if self.filepath == self.original_filepath:
            base, ext = os.path.splitext(self.original_filepath)
            new_path = f"{base}_edited{ext}"
            counter = 1
            while os.path.exists(new_path):
                new_path = f"{base}_edited_{counter}{ext}"
                counter += 1
            self.filepath = new_path
        try:
            self.workbook.save(self.filepath)
            self.modified = False
            self.file_label.config(text=self.filepath.split("/")[-1].split("\\")[-1], foreground="black")
            self.status_label.config(text="Tersimpan", foreground="green")
            self.statusbar.config(text=f"File disimpan: {self.filepath} (file asli tidak diubah)")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal menyimpan:\n{e}")

    def save_as(self):
        if not self.workbook:
            messagebox.showwarning("Peringatan", "Belum ada file yang dibuka")
            return
        path = filedialog.asksaveasfilename(
            title="Simpan Sebagai",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if not path:
            return
        try:
            self.workbook.save(path)
            self.filepath = path
            self.modified = False
            self.file_label.config(text=path.split("/")[-1].split("\\")[-1], foreground="black")
            self.status_label.config(text="Tersimpan", foreground="green")
            self.statusbar.config(text=f"File disimpan: {path}")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal menyimpan:\n{e}")

    def on_close(self):
        if self.modified:
            ans = messagebox.askyesnocancel("Simpan?", "Ada perubahan yang belum disimpan. Simpan dulu?")
            if ans is None:
                return
            if ans:
                self.save_file()
        self.root.destroy()

    # ------------------------------------------------------------------ SHEET TABS
    def _render_tabs(self):
        for w in self.tab_inner.winfo_children():
            w.destroy()
        if not self.workbook:
            return
        for name in self.workbook.sheetnames:
            btn = tk.Button(
                self.tab_inner, text=name, padx=12, pady=4,
                relief=tk.RAISED if name != self.current_sheet_name else tk.SUNKEN,
                bg="#2563eb" if name == self.current_sheet_name else "#f0f0f0",
                fg="white" if name == self.current_sheet_name else "black",
                font=("Segoe UI", 9, "bold" if name == self.current_sheet_name else "normal"),
                command=lambda n=name: self._switch_sheet(n)
            )
            btn.pack(side=tk.LEFT, padx=1, pady=2)

    def _switch_sheet(self, name):
        self.current_sheet_name = name
        self._render_tabs()
        self._render_table()

    # ------------------------------------------------------------------ TABLE
    def _render_table(self):
        self.tree.delete(*self.tree.get_children())

        if not self.workbook or not self.current_sheet_name:
            return

        ws = self.workbook[self.current_sheet_name]
        max_col = ws.max_column or 1
        max_row = ws.max_row or 1

        cols = ["#"] + [get_column_letter(c) for c in range(1, max_col + 1)]
        self.tree["columns"] = cols
        for c in cols:
            w = 50 if c == "#" else 140
            self.tree.heading(c, text=c)
            self.tree.column(c, width=w, minwidth=40, stretch=True)

        for row_idx in range(1, max_row + 1):
            values = [str(row_idx)]
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val = cell.value
                if val is None:
                    val = ""
                values.append(str(val))
            self.tree.insert("", tk.END, values=values, tags=(f"row_{row_idx}",))

        self.statusbar.config(text=f"Sheet: {self.current_sheet_name} | {max_row} baris x {max_col} kolom")

    # ------------------------------------------------------------------ LOAD DATA PASIEN
    def load_data_pasien(self):
        """Load patient data from .txt file or paste — only loads into memory."""
        win = tk.Toplevel(self.root)
        win.title("Load Data Pasien")
        win.geometry("750x650")
        win.transient(self.root)
        win.grab_set()

        ttk.Label(win, text="Load Data Pasien", font=("Segoe UI", 14, "bold")).pack(pady=(12, 4))
        ttk.Label(win, text=(
            "Paste data pasien di bawah, atau load dari file .txt.\n"
            "Format per baris: DD/MM/YYYY Nama No. Reg XXXXXXX dx: diagnosa\n\n"
            "Data akan disimpan di memori. Untuk proses data,\n"
            "gunakan tombol 'Proses Data' di toolbar atau menu Tools."
        ), justify=tk.LEFT, foreground="gray").pack(padx=16, anchor="w")

        # Load from file button
        load_frame = ttk.Frame(win)
        load_frame.pack(fill=tk.X, padx=16, pady=(4, 0))

        def load_from_txt():
            path = filedialog.askopenfilename(
                title="Pilih File Data Pasien",
                filetypes=[("Text files", "*.txt"), ("CSV files", "*.csv"), ("All files", "*.*")],
                parent=win
            )
            if not path:
                return
            try:
                with open(path, "r", encoding="utf-8") as f:
                    content = f.read()
            except UnicodeDecodeError:
                with open(path, "r", encoding="latin-1") as f:
                    content = f.read()
            text.delete("1.0", tk.END)
            text.insert("1.0", content)
            result_label.config(
                text=f"File dimuat: {path.split('/')[-1].split(chr(92))[-1]}",
                foreground="green"
            )

        ttk.Button(load_frame, text="Load dari File .txt", command=load_from_txt).pack(side=tk.LEFT)
        ttk.Label(load_frame, text="  atau paste langsung di bawah", foreground="gray").pack(side=tk.LEFT)

        text_frame = ttk.Frame(win)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=16, pady=8)

        text = tk.Text(text_frame, font=("Consolas", 10), wrap=tk.NONE)
        text_vsb = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=text.yview)
        text_hsb = ttk.Scrollbar(text_frame, orient=tk.HORIZONTAL, command=text.xview)
        text.configure(yscrollcommand=text_vsb.set, xscrollcommand=text_hsb.set)

        text.grid(row=0, column=0, sticky="nsew")
        text_vsb.grid(row=0, column=1, sticky="ns")
        text_hsb.grid(row=1, column=0, sticky="ew")
        text_frame.rowconfigure(0, weight=1)
        text_frame.columnconfigure(0, weight=1)

        # Pre-fill if data already loaded
        if self.patient_data:
            text.insert("1.0", self.patient_data)

        btn_frame = ttk.Frame(win)
        btn_frame.pack(fill=tk.X, padx=16, pady=(0, 12))

        result_label = ttk.Label(btn_frame, text="", foreground="blue")
        result_label.pack(side=tk.LEFT)

        def preview():
            raw = text.get("1.0", tk.END).strip()
            if not raw:
                result_label.config(text="Tidak ada data", foreground="red")
                return
            groups = self._parse_patient_data(raw)
            if not groups:
                result_label.config(text="Format data tidak dikenali", foreground="red")
                return
            dates = list(groups.keys())
            total = sum(len(v) for v in groups.values())
            sheet_info = f" Sheet tersedia: {len(self.workbook.sheetnames)}" if self.workbook else ""
            result_label.config(
                text=f"Ditemukan {len(dates)} tanggal, {total} pasien.{sheet_info}",
                foreground="blue"
            )

        def do_load():
            raw = text.get("1.0", tk.END).strip()
            if not raw:
                messagebox.showwarning("Peringatan", "Tidak ada data", parent=win)
                return
            groups = self._parse_patient_data(raw)
            if not groups:
                messagebox.showerror("Error", "Format data tidak dikenali", parent=win)
                return

            self.patient_data = raw
            dates = list(groups.keys())
            total = sum(len(v) for v in groups.values())

            self.statusbar.config(text=f"Data pasien dimuat: {len(dates)} tanggal, {total} pasien")
            self.status_label.config(text="Data dimuat", foreground="blue")

            messagebox.showinfo(
                "Sukses",
                f"Data pasien berhasil dimuat!\n"
                f"{len(dates)} tanggal, {total} pasien.\n\n"
                f"Untuk proses data, klik 'Proses Data' di toolbar.",
                parent=win
            )
            win.destroy()

        ttk.Button(btn_frame, text="Preview", command=preview).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text="Simpan Data", command=do_load).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text="Batal", command=win.destroy).pack(side=tk.RIGHT, padx=4)

    # ------------------------------------------------------------------ PROSES DATA
    # Template row constants (row indices in template sheet)
    TEMPLATE_DATA_START = 16   # first data row (Briefing)
    TEMPLATE_DATA_END = 22     # last data row (alat medis)
    TEMPLATE_ITEMS = 7         # number of items per table (rows 16-22)
    GAP_ROWS = 2               # empty rows between tables

    def proses_data(self):
        """Fill data pasien ke dalam sheet template, semua tanggal dalam 1 sheet."""
        if not self.workbook:
            messagebox.showwarning("Peringatan", "Buka file Excel terlebih dahulu")
            return
        if not self.patient_data:
            messagebox.showwarning(
                "Peringatan",
                "Belum ada data pasien yang dimuat.\n"
                "Klik 'Load Data Pasien' terlebih dahulu."
            )
            return

        groups = self._parse_patient_data(self.patient_data)
        if not groups:
            messagebox.showerror("Error", "Format data pasien tidak dikenali")
            return

        dates = list(groups.keys())
        total_patients = sum(len(v) for v in groups.values())

        template_ws = self.workbook.worksheets[0]
        template_name = template_ws.title

        msg = (
            f"Data: {len(dates)} tanggal, {total_patients} pasien\n"
            f"Template sheet: '{template_name}'\n\n"
            f"Semua tanggal akan ditulis dalam 1 sheet,\n"
            f"diisi data pasien otomatis, nomor urut lanjut.\n\n"
            f"Lanjutkan?"
        )
        if not messagebox.askyesno("Konfirmasi Proses", msg):
            return

        # Show progress window
        progress_win = tk.Toplevel(self.root)
        progress_win.title("Memproses Data...")
        progress_win.geometry("400x150")
        progress_win.transient(self.root)
        progress_win.grab_set()
        progress_win.resizable(False, False)

        ttk.Label(
            progress_win, text="Memproses data pasien...",
            font=("Segoe UI", 11, "bold")
        ).pack(pady=(16, 4))

        progress_label = ttk.Label(progress_win, text="Persiapan...", foreground="gray")
        progress_label.pack(pady=(0, 8))

        progress_bar = ttk.Progressbar(
            progress_win, orient=tk.HORIZONTAL, length=350, mode='determinate'
        )
        progress_bar.pack(padx=20, pady=(0, 16))
        progress_bar['maximum'] = len(dates)

        self.root.update_idletasks()

        def on_progress(date_idx, date_key):
            day, month, year = date_key
            progress_bar['value'] = date_idx + 1
            pct = int((date_idx + 1) / len(dates) * 100)
            progress_label.config(
                text=f"Tanggal {day} {BULAN_INDO[month]} ({date_idx + 1}/{len(dates)}) — {pct}%"
            )
            self.root.update_idletasks()

        self._fill_all_dates(template_ws, groups, on_progress)

        progress_win.destroy()

        self.modified = True
        self.status_label.config(text="Belum disimpan", foreground="orange")
        self.current_sheet_name = self.workbook.sheetnames[0]
        self._render_tabs()
        self._render_table()

        messagebox.showinfo(
            "Sukses",
            f"Berhasil memproses {len(dates)} tanggal dalam sheet '{template_ws.title}'!\n"
            f"Total {total_patients} pasien, nomor urut 1-{len(dates) * self.TEMPLATE_ITEMS}."
        )

    def _fill_all_dates(self, ws, groups, progress_callback=None):
        """Write all date blocks into a single sheet.

        Structure per date block (7 rows):
          - Row 1: Briefing (with date in col B)
          - Row 2: Timbang terima
          - Rows 3-6: Asuhan keperawatan etc. (with no CM)
          - Row 7: Alat medis (with EKG No. Reg)
        Between blocks: 2 empty rows.
        Numbering continues across blocks (1-7, 8-14, 15-21...).
        """
        from copy import copy as copy_style

        # Unmerge cells in the data/total area so we can write freely
        merged_to_remove = []
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_row >= self.TEMPLATE_DATA_START:
                merged_to_remove.append(merged_range)
        for mr in merged_to_remove:
            ws.unmerge_cells(str(mr))

        # Read template data rows (16-22) as reference
        template_rows = []
        max_col = ws.max_column or 1
        for src_row in range(self.TEMPLATE_DATA_START, self.TEMPLATE_DATA_END + 1):
            row_data = []
            for col in range(1, max_col + 1):
                cell = ws.cell(row=src_row, column=col)
                row_data.append({
                    'value': cell.value,
                    'font': copy_style(cell.font),
                    'border': copy_style(cell.border),
                    'fill': copy_style(cell.fill),
                    'alignment': copy_style(cell.alignment),
                    'number_format': cell.number_format,
                })
            template_rows.append(row_data)

        # Read template row heights
        template_heights = {}
        for src_row in range(self.TEMPLATE_DATA_START, self.TEMPLATE_DATA_END + 1):
            h = ws.row_dimensions[src_row].height
            template_heights[src_row - self.TEMPLATE_DATA_START] = h

        # First table starts at row 16 (template data start)
        current_row = self.TEMPLATE_DATA_START
        running_number = 1
        first_data_row = current_row
        dates = list(groups.keys())

        for date_idx, date_key in enumerate(dates):
            if progress_callback:
                progress_callback(date_idx, date_key)

            day, month, year = date_key
            patients = groups[date_key]
            date_obj = datetime(year, month, day)

            # Extract No. Reg and EKG info
            reg_numbers = []
            ekg_regs = []
            last_ekg_reg = None
            for patient_line in patients:
                reg_match = re.search(r'No\.\s*Reg\s+(\d+)', patient_line)
                if reg_match:
                    reg_num = reg_match.group(1)
                    has_ekg = 'ekg' in patient_line.lower()
                    reg_numbers.append((reg_num, has_ekg))
                    if has_ekg:
                        ekg_regs.append(reg_num)
                        last_ekg_reg = reg_num

            # Limit to 4 No. Reg: prioritize EKG patients, then take last non-EKG
            if len(reg_numbers) > 4:
                ekg_items = [(r, e) for r, e in reg_numbers if e]
                non_ekg_items = [(r, e) for r, e in reg_numbers if not e]
                remaining = 4 - len(ekg_items)
                if remaining > 0:
                    selected = ekg_items + non_ekg_items[-remaining:]
                else:
                    selected = ekg_items[-4:]
                selected_regs = [r for r, _ in selected]
            else:
                selected_regs = [r for r, _ in reg_numbers]
            reg_list_str = ", ".join(selected_regs)

            # Write 7 data rows for this date
            for item_idx in range(self.TEMPLATE_ITEMS):
                dest_row = current_row + item_idx
                tmpl = template_rows[item_idx]

                # Set row height from template
                if item_idx in template_heights and template_heights[item_idx]:
                    ws.row_dimensions[dest_row].height = template_heights[item_idx]

                for col in range(1, max_col + 1):
                    dest_cell = ws.cell(row=dest_row, column=col)
                    src_data = tmpl[col - 1]

                    # Copy formatting
                    dest_cell.font = copy_style(src_data['font'])
                    dest_cell.border = copy_style(src_data['border'])
                    dest_cell.fill = copy_style(src_data['fill'])
                    dest_cell.alignment = copy_style(src_data['alignment'])
                    dest_cell.number_format = src_data['number_format']

                    val = src_data['value']

                    # Column A: sequential number
                    if col == 1:
                        dest_cell.value = running_number + item_idx
                        continue

                    # Column B: date on first row (Briefing), empty on others
                    if col == 2:
                        if item_idx == 0:
                            dest_cell.value = date_obj
                            dest_cell.number_format = 'DD/MM/YYYY'
                        else:
                            dest_cell.value = None
                        continue

                    # Column G: formula =SUM(E*F) adjusted for current row
                    if col == 7 and isinstance(val, str) and val.startswith('='):
                        dest_cell.value = f"=SUM(E{dest_row}*F{dest_row})"
                        continue

                    # Column D: fill "no CM" and "alat medis" text
                    if col == 4 and isinstance(val, str):
                        if val.rstrip().endswith("no CM"):
                            dest_cell.value = f"{val} {reg_list_str}."
                            continue
                        if "agar siap pakai dengan" in val.lower():
                            if last_ekg_reg:
                                dest_cell.value = f"{val.rstrip()} {last_ekg_reg}."
                            else:
                                dest_cell.value = val
                            continue

                    # Column J: only copy on first date block
                    if col >= 10 and date_idx > 0:
                        dest_cell.value = None
                        continue

                    # Default: copy value as-is
                    dest_cell.value = val

            running_number += self.TEMPLATE_ITEMS
            current_row += self.TEMPLATE_ITEMS

            # Add gap rows between tables (not after last)
            if date_idx < len(dates) - 1:
                current_row += self.GAP_ROWS

        # Write Total row after all data
        last_data_row = current_row - 1
        total_row = current_row + 1
        ws.cell(row=total_row, column=1).value = "Total "
        total_formula = f"=SUM(G{first_data_row}:G{last_data_row})"
        ws.cell(row=total_row, column=7).value = total_formula

        # Copy Total row formatting from template row 31
        for col in range(1, max_col + 1):
            src_cell = ws.cell(row=31, column=col)
            if total_row != 31:
                dest_cell = ws.cell(row=total_row, column=col)
                dest_cell.font = copy_style(src_cell.font)
                dest_cell.border = copy_style(src_cell.border)
                dest_cell.fill = copy_style(src_cell.fill)
                dest_cell.alignment = copy_style(src_cell.alignment)

        # Write NB row
        nb_row = total_row + 2
        ws.cell(row=nb_row, column=1).value = "NB"
        ws.cell(row=nb_row, column=2).value = "1 Perawat 4 Pasien"

        # Write formula row
        formula_row = nb_row + 1
        ws.cell(row=formula_row, column=7).value = f"=G{total_row}"
        ws.cell(row=formula_row, column=8).value = 20
        ws.cell(row=formula_row, column=9).value = f"=G{formula_row}*H{formula_row}"
        ws.cell(row=formula_row, column=10).value = f"=I{formula_row}/60"

        # Clear old template rows that are below our data (rows 23-34 from original)
        # Only needed if our data ends before the old template area
        clear_start = max(formula_row + 1, self.TEMPLATE_DATA_END + 1)
        for row in range(clear_start, 35):
            for col in range(1, max_col + 1):
                ws.cell(row=row, column=col).value = None

    def _parse_patient_data(self, raw_text):
        """Parse patient data text and group by date.
        Returns OrderedDict: {(day, month, year): [lines...]}
        """
        groups = OrderedDict()
        pattern = re.compile(
            r'(\d{1,2})[/\-](\d{1,2})[/\-](\d{2,4})\s+(.+)'
        )

        for line in raw_text.strip().splitlines():
            line = line.strip()
            if not line:
                continue
            m = pattern.match(line)
            if m:
                day = int(m.group(1))
                month = int(m.group(2))
                year = int(m.group(3))
                rest = m.group(4).strip()
                key = (day, month, year)
                if key not in groups:
                    groups[key] = []
                groups[key].append(rest)

        return groups

    def _copy_images(self, source_ws, target_ws):
        """Copy all images from source worksheet to target worksheet."""
        for img in source_ws._images:
            try:
                img.ref.seek(0)
                img_data = img.ref.read()
                img.ref.seek(0)
            except Exception:
                try:
                    img_data = img._data()
                except Exception:
                    continue

            new_img = XlImage(io.BytesIO(img_data))
            new_img.width = img.width
            new_img.height = img.height

            anchor = img.anchor
            if isinstance(anchor, TwoCellAnchor):
                new_anchor = TwoCellAnchor()
                new_anchor._from = AnchorMarker(
                    col=anchor._from.col,
                    colOff=anchor._from.colOff,
                    row=anchor._from.row,
                    rowOff=anchor._from.rowOff
                )
                new_anchor.to = AnchorMarker(
                    col=anchor.to.col,
                    colOff=anchor.to.colOff,
                    row=anchor.to.row,
                    rowOff=anchor.to.rowOff
                )
                new_img.anchor = new_anchor
            elif isinstance(anchor, OneCellAnchor):
                new_anchor = OneCellAnchor()
                new_anchor._from = AnchorMarker(
                    col=anchor._from.col,
                    colOff=anchor._from.colOff,
                    row=anchor._from.row,
                    rowOff=anchor._from.rowOff
                )
                new_img.anchor = new_anchor

            target_ws.add_image(new_img)


def main():
    root = tk.Tk()

    # Style
    style = ttk.Style()
    style.theme_use("clam")
    style.configure("Treeview", rowheight=28, font=("Segoe UI", 10))
    style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))

    app = ExcelEditorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
